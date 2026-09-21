"""Excel rendering for the weekly sales report.

Produces the workbook that is both downloaded from the app and attached to the
weekly email. `scripts/reports/generate_periodic_report.py` renders through this
same code, so the file a scheduled run emails is byte-for-byte the file the CLI
has been producing since Phase 1.

Money is displayed in MILLIONS of VND. Each comma at the tail of an Excel number
format divides the DISPLAYED value by 1000, so two commas render 159,745,427,865
as "159,745" — rounded to the nearest million for display only. The cell still
holds the exact Decimal, so sub-million remainders are preserved and any formula
over these cells stays exact. The `Raw Data` sheet deliberately stays in whole
VND: its day x store x department rows are often below a million, which would
render as "0".
"""
from io import BytesIO

from .aggregate import (LOWER_IS_BETTER, METRICS, ZERO_METRICS, aggregate,
                        delta)
from .period import week_label

MONEY_FORMAT = '#,##0,,'        # whole millions -- the aggregate totals
MONEY_1DP_FORMAT = '#,##0.0,,'  # millions to 1dp -- the per-unit averages
COUNT_FORMAT = '#,##0'
PCT_FORMAT = '#,##0.0"%"'
POINT_FORMAT = '#,##0.0"pp"'

# (label, metric key, unit).
#
# The two per-unit averages use `money1` rather than `money`: they sit around
# 12-36 million, so rounding them to whole millions renders an average unit
# price of 12,286,107 as a bare "12" and makes a move from 12.2m to 12.4m
# invisible. One decimal keeps the same millions unit and restores that signal.
# The aggregate totals are in the tens of thousands of millions, where a decimal
# would be noise.
#
# Labels are kept short so the columns can be narrow enough to fit all three
# blocks on one screen. Money labels MUST keep the "(VND m)" suffix — the delta
# block strips it, and the unit has to stay visible on the two value blocks.
#
# Order and units are pinned to `aggregate.METRICS`; a test asserts they agree,
# because the sheet addresses metric *i* of a block by position.
_LABELS = {
    'total_sales':           'Sales (VND m)',
    'qty_sold':              'Qty sold',
    'bills':                 'Bills',
    'avg_unit_price':        'Avg price (VND m)',
    'avg_transaction_value': 'Avg txn (VND m)',
    'avg_discount_pct':      'Discount %',
    'returns_value':         'Returns (VND m)',
}

METRIC_COLUMNS = [(_LABELS[key], key, unit) for key, unit in METRICS]

# Only unfavourable moves are coloured; favourable ones keep the default text
# colour, so the eye goes straight to the problems.
BAD_COLOUR = 'C00000'

_UNIT_FORMAT = {
    'money': MONEY_FORMAT,
    'money1': MONEY_1DP_FORMAT,
    'count': COUNT_FORMAT,
    'pct': PCT_FORMAT,
}

# The sheets the workbook carries, in order. Deliberately NOT PERIOD_LABELS:
# WTD is an app-only view, so the emailed attachment is unchanged by it.
SHEET_NAMES = {'WOW': 'Week by Week', 'MTD': 'MTD', 'YTD': 'YTD'}


def _write_period_sheet(ws, label, windows, cur_data, prior_data):
    from openpyxl.styles import Alignment, Font, PatternFill

    bold = Font(bold=True)
    total_fill = PatternFill('solid', fgColor='FFF2CC')
    grand_fill = PatternFill('solid', fgColor='E2EFDA')

    (c_from, c_to), (p_from, p_to) = windows

    cur_dept, cur_store, cur_grand, store_names = cur_data
    pri_dept, pri_store, pri_grand, pri_names = prior_data
    store_names = {**pri_names, **store_names}

    if label == 'WOW':
        title = 'Week by Week'
        headline = f'{week_label(c_from)}  vs  {week_label(p_from)}'
        alignment_note = ('Latest COMPLETE week against the week before it (the '
                          'current partial week is excluded). Both sides are whole '
                          'weeks on the same weekdays. This sheet is week-over-week, '
                          'NOT year-over-year — MTD and YTD compare against last year.')
    else:
        title = label
        headline = f'{c_from} to {c_to}  vs  {p_from} to {p_to}'
        alignment_note = 'Prior-year window is calendar-date aligned.'

    ws['A1'] = f'{title} — {headline}'
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = (f'{c_from} ({c_from:%a}) to {c_to} ({c_to:%a})   vs   '
                f'{p_from} ({p_from:%a}) to {p_to} ({p_to:%a})')
    ws['A2'].font = Font(size=10)
    legend = '   '.join(f'{code} = {store_names[code]}'
                        for code in sorted(store_names))
    ws['A4'] = f'Stores:   {legend}'
    ws['A4'].font = Font(italic=True, size=9)

    ws['A3'] = (f'{alignment_note} Money columns are shown in MILLIONS of VND '
                '(display rounding only — cells hold exact values). Department bill '
                'counts count bills CONTAINING that department and do not sum to the '
                'store total; store/grand totals are document-grain.')
    ws['A3'].font = Font(italic=True, size=9)

    # Two header rows. The sheet is laid out in three BLOCKS -- every metric for
    # the current period side by side, then the same metrics for the prior
    # period, then the deltas -- rather than interleaving the views of each
    # metric. Reading one period straight across is the common case; comparing a
    # single metric across periods is the exception.
    top, sub = 5, 6
    ws.cell(top, 1, 'Store').font = bold
    ws.cell(top, 2, 'Department').font = bold
    ws.merge_cells(start_row=top, start_column=1, end_row=sub, end_column=1)
    ws.merge_cells(start_row=top, start_column=2, end_row=sub, end_column=2)

    if label == 'WOW':
        cur_block = f'Latest week — {week_label(c_from)}'
        pri_block = f'Prior week — {week_label(p_from)}'
    else:
        cur_block = f'Current — {c_from} to {c_to}'
        pri_block = f'Prior year — {p_from} to {p_to}'

    width = len(METRIC_COLUMNS)
    blocks = [
        (cur_block, PatternFill('solid', fgColor='DDEBF7')),
        (pri_block, PatternFill('solid', fgColor='EDEDED')),
        ('Δ %   (red = unfavourable)', PatternFill('solid', fgColor='FFF2CC')),
    ]

    for b, (block_label, fill) in enumerate(blocks):
        start = 3 + b * width
        ws.merge_cells(start_row=top, start_column=start,
                       end_row=top, end_column=start + width - 1)
        c = ws.cell(top, start, block_label)
        c.font = bold
        c.fill = fill
        c.alignment = Alignment(horizontal='center')
        is_delta_block = b == len(blocks) - 1
        for i, (metric_label, _key, unit) in enumerate(METRIC_COLUMNS):
            # The delta block is already all-percentages; repeating "(VND m)"
            # there would misdescribe it.
            text = metric_label.replace(' (VND m)', '') if is_delta_block else metric_label
            sc = ws.cell(sub, start + i, text)
            sc.font = bold
            sc.fill = fill
            sc.alignment = Alignment(horizontal='center', wrap_text=True)

    def emit(row, store_label, dept_label, cur_m, pri_m, fill=None):
        a = ws.cell(row, 1, store_label)
        b = ws.cell(row, 2, dept_label)
        if fill:
            a.fill = b.fill = fill
            a.font = b.font = bold
        cur_m = cur_m or ZERO_METRICS
        pri_m = pri_m or ZERO_METRICS
        for i, (_lbl, key, unit) in enumerate(METRIC_COLUMNS):
            cv = cur_m.get(key)
            pv = pri_m.get(key)
            diff, pct = delta(cv, pv)
            fmt = _UNIT_FORMAT[unit]
            # For a percentage metric the delta is a percentage-POINT gap, not a
            # ratio of ratios -- a "26.4% vs 26.5%" change is +0.1pp, and
            # expressing that as +0.4% would invite the wrong reading.
            cells = (
                (cv, fmt),
                (pv, fmt),
                (diff, POINT_FORMAT) if unit == 'pct' else (pct, PCT_FORMAT),
            )
            for b, (val, vfmt) in enumerate(cells):
                cell = ws.cell(row, 3 + b * len(METRIC_COLUMNS) + i, val)
                cell.number_format = vfmt
                # Red only in the delta block, and only for a move in the
                # unfavourable direction for THAT metric.
                bad = False
                if b == len(blocks) - 1 and val is not None and val != 0:
                    bad = (val > 0) if key in LOWER_IS_BETTER else (val < 0)
                if fill:
                    cell.fill = fill
                    cell.font = Font(bold=True, color=BAD_COLOUR) if bad else bold
                elif bad:
                    cell.font = Font(color=BAD_COLOUR)
        return row + 1

    # Two-level Excel outline:
    #   level 1 (the "1" button, and how the file opens) -- store TOTAL rows
    #                                                       plus the grand total
    #   level 2 (the "2" button)                         -- every department row
    # Department rows sit ABOVE their store's TOTAL, so summaryBelow must stay
    # True (openpyxl's default) for the +/- control to attach to the TOTAL row.
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.sheet_properties.outlinePr.applyStyles = False

    row = sub + 1
    all_stores = sorted(set(cur_store) | set(pri_store))
    for store in all_stores:
        depts = sorted({d for (s, d) in cur_dept if s == store} |
                       {d for (s, d) in pri_dept if s == store})
        first_detail = row
        for dept in depts:
            row = emit(row, store, dept,
                       cur_dept.get((store, dept)), pri_dept.get((store, dept)))
        last_detail = row - 1
        total_row = row
        row = emit(row, store, 'TOTAL',
                   cur_store.get(store), pri_store.get(store), total_fill)

        # Group this store's department rows and start them collapsed, so the
        # sheet opens on the store view. `collapsed` belongs on the summary row
        # (the TOTAL), not on the hidden detail rows.
        for r in range(first_detail, last_detail + 1):
            ws.row_dimensions[r].outlineLevel = 1
            ws.row_dimensions[r].hidden = True
        if depts:
            ws.row_dimensions[total_row].collapsed = True

        row += 1   # spacer between stores, deliberately left at level 0

    emit(row, 'ALL STORES', 'TOTAL', cur_grand, pri_grand, grand_fill)

    from openpyxl.utils import get_column_letter

    # Sized so ALL THREE blocks (latest, prior, delta) fit on screen without
    # horizontal scrolling. Widths are per metric and measured from rendered
    # content rather than set uniformly: the widest value actually produced is
    # 7 chars for money, 6 for qty, 5 for bills, and 8 in the delta block
    # ("1,050.0%"). Each width carries ~2 chars of headroom on top of that so
    # larger figures in future periods do not clip.
    #   px = round(width * 7) + 5
    #   A(11) + B(10) + 2 value blocks(455px each) + delta block(476px)
    #   = 82 + 75 + 910 + 476 = 1543px at 100% zoom
    # Reference point: a 15" MacBook at default scaling gives ~1380px of grid,
    # which fits the two value blocks; the wider layouts (16" / "More Space")
    # give ~1700px+ and fit the delta block too.
    STORE_W, DEPT_W, ZOOM = 11, 10, 100
    VALUE_W = {
        'total_sales':           10,
        'qty_sold':               8,
        'bills':                  7,
        'avg_unit_price':         9,
        'avg_transaction_value':  9,
        'avg_discount_pct':       8,
        'returns_value':          9,
    }
    DELTA_W = 9   # the whole delta block is percentages -- one width suffices

    ws.freeze_panes = 'C7'
    ws.column_dimensions['A'].width = STORE_W
    ws.column_dimensions['B'].width = DEPT_W
    for b in range(len(blocks)):
        is_delta = b == len(blocks) - 1
        for i, (_lbl, key, _unit) in enumerate(METRIC_COLUMNS):
            col_letter = get_column_letter(3 + b * width + i)
            ws.column_dimensions[col_letter].width = (
                DELTA_W if is_delta else VALUE_W[key])
    # The prior-period block opens COLLAPSED behind an outline button. All
    # three blocks together are ~1543px, which overflows a 15" screen; the
    # prior figures are reference material, while the current figures and the
    # deltas are what actually gets read. Collapsing them puts the delta block
    # on screen next to the current one, and a single click brings them back.
    #
    # Deliberately NOT `ws.column_dimensions.group()`. That helper DELETES the
    # per-column dimensions across the range and leaves a single spanning one
    # keyed to the first column, so the widths set just above are lost: K..P
    # fall back to Excel's default 13 and the block silently gets wider on
    # expand than it was designed to be. Measured, not assumed --
    # group('J','P') on widths [10,8,7,9,9,8,9] yields [10,13,13,13,13,13,13].
    # Setting the outline attributes per column preserves them.
    prior_start = 3 + width
    for i in range(width):
        dim = ws.column_dimensions[get_column_letter(prior_start + i)]
        dim.outlineLevel = 1
        dim.hidden = True
    # Excel draws the +/- button on the column AFTER the group and marks THAT
    # column collapsed -- the same summary-cell rule the row outline above
    # follows, mirrored to the right. summaryRight is openpyxl's default; it is
    # set explicitly because the button lands on the delta block, and having
    # that be accidental would be easy to break.
    ws.sheet_properties.outlinePr.summaryRight = True
    ws.column_dimensions[get_column_letter(prior_start + width)].collapsed = True

    # Narrow columns push the wrapped headers onto two or three lines.
    ws.row_dimensions[sub].height = 44
    ws.sheet_view.zoomScale = ZOOM


def _write_raw_sheet(ws, dept_rows):
    from openpyxl.styles import Font

    # Whole VND here, not millions: a single day x store x department row is
    # often under a million and would round to "0".
    headers = ['Day', 'Store code', 'Store name', 'Department',
               'Sale net (VND, ex-VAT)', 'Return net (VND, ex-VAT)',
               'Sale gross (VND, list, ex-VAT)',
               'Qty sold', 'Qty returned', 'Bills containing dept']
    for i, h in enumerate(headers, start=1):
        c = ws.cell(1, i, h)
        c.font = Font(bold=True)

    for r, row in enumerate(sorted(dept_rows,
                                   key=lambda x: (x['day'], x['store_code'], x['department'])),
                            start=2):
        ws.cell(r, 1, row['day']).number_format = 'yyyy-mm-dd'
        ws.cell(r, 2, row['store_code'])
        ws.cell(r, 3, row['store_name'])
        ws.cell(r, 4, row['department'])
        for i, key in enumerate(('sale_net', 'return_net', 'sale_gross'), start=5):
            ws.cell(r, i, row[key]).number_format = '#,##0'
        ws.cell(r, 8, row['qty_sold']).number_format = '#,##0'
        ws.cell(r, 9, row['qty_returned']).number_format = '#,##0'
        ws.cell(r, 10, row['dept_bills']).number_format = '#,##0'

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip('ABCDEFGHIJ', (12, 11, 26, 13, 20, 20, 24, 11, 13, 21)):
        ws.column_dimensions[col].width = width


def build_workbook_object(windows, dept_rows, store_rows):
    """Render the three period sheets plus `Raw Data` into a Workbook."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)

    # Iterates the SHEETS, not every period: the payload carries WTD as well,
    # and the workbook deliberately does not.
    for label, sheet_name in SHEET_NAMES.items():
        cur_win, prior_win = windows[label]
        ws = wb.create_sheet(sheet_name)
        _write_period_sheet(
            ws, label, windows[label],
            aggregate(dept_rows, store_rows, cur_win),
            aggregate(dept_rows, store_rows, prior_win),
        )

    _write_raw_sheet(wb.create_sheet('Raw Data'), dept_rows)
    return wb


def build_workbook(path, as_of, windows, dept_rows, store_rows):
    """Render and save to `path`. Used by the CLI generator."""
    wb = build_workbook_object(windows, dept_rows, store_rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def build_workbook_bytes(windows, dept_rows, store_rows) -> bytes:
    """Render to bytes. Used by the scheduled run, which stores the workbook in
    Postgres — `deploy.yml` deletes `document/` on the server, so anything
    written to disk there would not survive the next release."""
    buf = BytesIO()
    build_workbook_object(windows, dept_rows, store_rows).save(buf)
    return buf.getvalue()
