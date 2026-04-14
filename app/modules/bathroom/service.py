"""
Bathroom Products Service
Manages product catalog data for bathroom fixture brands (Dolomite, Valsir, Paffoni).
"""
import csv
import json
import os
from typing import Dict, Any, List, Optional

from app.core.database import get_postgres_connection


# Brand name normalization: Excel sheet names → standard brand keys
BRAND_SHEET_MAP = {
    'dolomite': 'dolomite',
    'valsir': 'valsir',
    'paffoni': 'paffoni',
}

# Brands where shorten_code grouping is meaningful (have color/finish variants)
BRANDS_WITH_VARIANTS = {'dolomite', 'paffoni'}


class BathroomService:

    def init_database(self) -> Dict[str, Any]:
        conn = None
        try:
            conn = get_postgres_connection()
            if not conn:
                return {'success': False, 'error': 'Failed to connect to database'}

            cursor = conn.cursor()

            # Sequence for product IDs (prefix 5)
            cursor.execute("""
                CREATE SEQUENCE IF NOT EXISTS bathroom_product_id_seq
                    START WITH 500000001 INCREMENT BY 1 NO MAXVALUE NO CYCLE
            """)

            # Products table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS bathroom_products (
                    id BIGINT PRIMARY KEY DEFAULT nextval('bathroom_product_id_seq'),
                    brand VARCHAR(50) NOT NULL,
                    catalog_year INTEGER,
                    model_code VARCHAR(100) NOT NULL,
                    shorten_code VARCHAR(20),
                    category VARCHAR(100),
                    description TEXT,
                    material VARCHAR(100),
                    base_price_eur NUMERIC(12,2) NOT NULL,
                    warranty VARCHAR(100),
                    section VARCHAR(100),
                    collection VARCHAR(100),
                    finish VARCHAR(100),
                    specs JSONB,
                    is_active BOOLEAN DEFAULT TRUE,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)

            # Unique constraint for upsert
            cursor.execute("""
                CREATE UNIQUE INDEX IF NOT EXISTS idx_bp_brand_model
                    ON bathroom_products(brand, model_code)
            """)

            # Search indexes
            cursor.execute("""
                CREATE INDEX IF NOT EXISTS idx_bp_model_code
                    ON bathroom_products(model_code)
            """)
            cursor.execute("""
                CREATE INDEX IF NOT EXISTS idx_bp_shorten_code
                    ON bathroom_products(shorten_code)
            """)
            cursor.execute("""
                CREATE INDEX IF NOT EXISTS idx_bp_brand_active
                    ON bathroom_products(brand, is_active)
            """)
            cursor.execute("""
                CREATE INDEX IF NOT EXISTS idx_bp_category
                    ON bathroom_products(category)
            """)
            cursor.execute("""
                CREATE INDEX IF NOT EXISTS idx_bp_specs
                    ON bathroom_products USING GIN (specs)
            """)

            conn.commit()
            cursor.close()
            return {'success': True, 'message': 'Bathroom products tables initialized'}

        except Exception as e:
            if conn:
                conn.rollback()
            return {'success': False, 'error': str(e)}
        finally:
            if conn:
                conn.close()

    def import_from_excel(self, file_path: str) -> Dict[str, Any]:
        try:
            import openpyxl
        except ImportError:
            return {'success': False, 'error': 'openpyxl is not installed'}

        if not os.path.exists(file_path):
            return {'success': False, 'error': f'File not found: {file_path}'}

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except Exception as e:
            return {'success': False, 'error': f'Failed to open Excel file: {e}'}

        results = {}
        for sheet_name, brand_key in BRAND_SHEET_MAP.items():
            if sheet_name not in wb.sheetnames:
                results[brand_key] = {'error': f'Sheet "{sheet_name}" not found'}
                continue

            ws = wb[sheet_name]
            rows = self._read_excel_sheet(ws, brand_key)
            result = self._upsert_products(brand_key, rows)
            results[brand_key] = result

        return {'success': True, 'data': results}

    def _read_excel_sheet(self, ws, brand_key: str) -> List[Dict]:
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        headers = [h.strip().lower() if h else '' for h in headers]

        rows = []
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
            row_dict = dict(zip(headers, row))

            model_code = str(row_dict.get('model_code', '') or '').strip()
            if not model_code:
                continue

            # Normalize model_code: uppercase, remove spaces
            model_code = model_code.upper().replace(' ', '')

            # shorten_code: first 5 chars for brands with variants, NULL for others
            shorten_code = None
            if brand_key in BRANDS_WITH_VARIANTS:
                shorten_code = model_code[:5] if len(model_code) >= 5 else model_code

            price = row_dict.get('price')
            if price is None:
                continue
            try:
                price = float(price)
            except (ValueError, TypeError):
                continue

            # Map Excel 'color' to 'finish'
            finish = str(row_dict.get('color', '') or '').strip() or None

            rows.append({
                'brand': brand_key,
                'catalog_year': row_dict.get('catalog_year'),
                'model_code': model_code,
                'shorten_code': shorten_code,
                'category': None,
                'description': str(row_dict.get('description', '') or '').strip() or None,
                'material': str(row_dict.get('material', '') or '').strip() or None,
                'base_price_eur': price,
                'warranty': str(row_dict.get('warranty', '') or '').strip() or None,
                'section': str(row_dict.get('section', '') or '').strip() or None,
                'collection': str(row_dict.get('collection', '') or '').strip() or None,
                'finish': finish,
                'specs': None,
            })

        return rows

    def import_from_csv(self, file_path: str, brand: str) -> Dict[str, Any]:
        if not os.path.exists(file_path):
            return {'success': False, 'error': f'File not found: {file_path}'}

        brand_key = brand.lower().strip()

        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                rows = []
                for row in reader:
                    model_code = (row.get('model_code') or '').strip().upper().replace(' ', '')
                    if not model_code:
                        continue

                    price = row.get('price') or row.get('base_price_eur')
                    if not price:
                        continue
                    try:
                        price = float(price)
                    except (ValueError, TypeError):
                        continue

                    shorten_code = None
                    if brand_key in BRANDS_WITH_VARIANTS:
                        shorten_code = model_code[:5] if len(model_code) >= 5 else model_code

                    # Parse specs JSON if present
                    specs = None
                    specs_raw = (row.get('specs') or '').strip()
                    if specs_raw:
                        try:
                            specs = json.loads(specs_raw)
                        except json.JSONDecodeError:
                            pass

                    rows.append({
                        'brand': brand_key,
                        'catalog_year': int(row['catalog_year']) if row.get('catalog_year') else None,
                        'model_code': model_code,
                        'shorten_code': shorten_code,
                        'category': (row.get('category') or '').strip() or None,
                        'description': (row.get('description') or '').strip() or None,
                        'material': (row.get('material') or '').strip() or None,
                        'base_price_eur': price,
                        'warranty': (row.get('warranty') or '').strip() or None,
                        'section': (row.get('section') or '').strip() or None,
                        'collection': (row.get('collection') or '').strip() or None,
                        'finish': (row.get('finish') or row.get('color') or '').strip() or None,
                        'specs': json.dumps(specs) if specs else None,
                    })
        except Exception as e:
            return {'success': False, 'error': f'Failed to read CSV: {e}'}

        return self._upsert_products(brand_key, rows)

    def _upsert_products(self, brand_key: str, rows: List[Dict]) -> Dict[str, Any]:
        if not rows:
            return {'success': False, 'error': 'No valid rows to import'}

        conn = None
        try:
            conn = get_postgres_connection()
            if not conn:
                return {'success': False, 'error': 'Failed to connect to database'}

            cursor = conn.cursor()

            inserted = 0
            updated = 0
            model_codes = []

            for row in rows:
                model_codes.append(row['model_code'])

                cursor.execute("""
                    INSERT INTO bathroom_products
                        (brand, catalog_year, model_code, shorten_code, category,
                         description, material, base_price_eur, warranty,
                         section, collection, finish, specs,
                         is_active, created_at, updated_at)
                    VALUES
                        (%(brand)s, %(catalog_year)s, %(model_code)s, %(shorten_code)s,
                         %(category)s, %(description)s, %(material)s, %(base_price_eur)s,
                         %(warranty)s, %(section)s, %(collection)s, %(finish)s,
                         %(specs)s,
                         TRUE, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
                    ON CONFLICT (brand, model_code) DO UPDATE SET
                        catalog_year = EXCLUDED.catalog_year,
                        shorten_code = EXCLUDED.shorten_code,
                        category = COALESCE(EXCLUDED.category, bathroom_products.category),
                        description = COALESCE(EXCLUDED.description, bathroom_products.description),
                        material = COALESCE(EXCLUDED.material, bathroom_products.material),
                        base_price_eur = EXCLUDED.base_price_eur,
                        warranty = COALESCE(EXCLUDED.warranty, bathroom_products.warranty),
                        section = COALESCE(EXCLUDED.section, bathroom_products.section),
                        collection = COALESCE(EXCLUDED.collection, bathroom_products.collection),
                        finish = COALESCE(EXCLUDED.finish, bathroom_products.finish),
                        specs = COALESCE(EXCLUDED.specs, bathroom_products.specs),
                        is_active = TRUE,
                        updated_at = CURRENT_TIMESTAMP
                    RETURNING (xmax = 0) AS is_insert
                """, row)

                result = cursor.fetchone()
                if result and result[0]:
                    inserted += 1
                else:
                    updated += 1

            # Mark products not in this import as inactive
            if model_codes:
                cursor.execute("""
                    UPDATE bathroom_products
                    SET is_active = FALSE, updated_at = CURRENT_TIMESTAMP
                    WHERE brand = %s
                      AND model_code != ALL(%s)
                      AND is_active = TRUE
                """, (brand_key, model_codes))
                deactivated = cursor.rowcount
            else:
                deactivated = 0

            conn.commit()
            cursor.close()

            return {
                'success': True,
                'brand': brand_key,
                'inserted': inserted,
                'updated': updated,
                'deactivated': deactivated,
                'total': len(rows)
            }

        except Exception as e:
            if conn:
                conn.rollback()
            return {'success': False, 'error': str(e)}
        finally:
            if conn:
                conn.close()
