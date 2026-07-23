"""
Master-data Excel report: Product master list.

One row per distinct item (INVN_SBS_ITEM.SID) that was received via
voucher between --start-date and --end-date. Same import cohort the
Sale and COGS report uses.

Columns:
  upc                 — TO_CHAR(INVN_SBS_ITEM.UPC)
  product_code        — INVN_SBS_ITEM.ALU (e.g. MAR-SS15-WRTW-TSH)
  product_name        — INVN_SBS_ITEM.DESCRIPTION1
  product_description — INVN_SBS_ITEM.TEXT1
  brand_code          — VENDOR.VEND_CODE
  brand_name          — VENDOR.VEND_NAME
  department          — DCS.D_LONG_NAME
  category            — DCS.C_LONG_NAME
  item_size           — INVN_SBS_ITEM.ITEM_SIZE
  season              — INVN_SBS_ITEM.UDF5_STRING

Output: document/reports/product_master_list_<start>_<today>.xlsx
"""
import argparse
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from app.core.database import get_oracle_connection


QUERY = r"""
    WITH period_received AS (
        SELECT DISTINCT vi.item_sid AS item_sid
        FROM VOUCHER vh
        JOIN VOU_ITEM vi ON vi.vou_sid = vh.SID
        WHERE vh.vou_class = 0
          AND vh.status = 4
          AND vh.vou_type IN (0, 1)
          AND vh.post_date >= TO_DATE(:start_date, 'YYYY-MM-DD')
          AND TRUNC(vh.post_date) <= TO_DATE(:end_date, 'YYYY-MM-DD')
    )
    SELECT
        TO_CHAR(i.UPC)          AS upc,
        i.ALU                   AS product_code,
        i.DESCRIPTION1          AS product_name,
        i.TEXT1                 AS product_description,
        v.VEND_CODE             AS brand_code,
        v.VEND_NAME             AS brand_name,
        dep.D_LONG_NAME         AS department,
        dep.C_LONG_NAME         AS category,
        i.ITEM_SIZE             AS item_size,
        i.UDF5_STRING           AS season
    FROM INVN_SBS_ITEM i
    JOIN period_received pr ON pr.item_sid = i.SID
    LEFT JOIN VENDOR v ON v.SID = i.VEND_SID
    LEFT JOIN DCS dep  ON dep.SID = i.DCS_SID
    ORDER BY brand_code, department, category, product_code, item_size
"""


def fetch_rows(start_date: str, end_date: str):
    conn = get_oracle_connection()
    if conn is None:
        raise RuntimeError('Could not open Oracle connection')
    try:
        cur = conn.cursor()
        try:
            cur.execute(QUERY, {'start_date': start_date, 'end_date': end_date})
            cols = [d[0].lower() for d in cur.description]
            rows = cur.fetchall()
        finally:
            cur.close()
    finally:
        conn.close()
    return [dict(zip(cols, r)) for r in rows]


def shape_rows(raw_rows):
    out = []
    for r in raw_rows:
        out.append({k: (r.get(k) or '') for k in [
            'upc', 'product_code', 'product_name', 'product_description',
            'brand_code', 'brand_name', 'department', 'category',
            'item_size', 'season',
        ]})
    return out


def write_xlsx(rows, out_path: Path, start_date: str, end_date: str, as_of: date):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Product Master List'

    headers = [
        'upc', 'product_code', 'product_name', 'product_description',
        'brand_code', 'brand_name', 'department', 'category',
        'item_size', 'season',
    ]
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='305496', end_color='305496', fill_type='solid')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for r in rows:
        ws.append([r[h] for h in headers])

    # Autosize (capped for readability)
    for col_idx, header in enumerate(headers, start=1):
        cells = [ws.cell(row=r, column=col_idx).value
                 for r in range(2, ws.max_row + 1)] if ws.max_row > 1 else []
        max_len = max(len(str(header)),
                      *(len(str(v or '')) for v in cells)) if cells else len(header)
        # Wider cap for description column
        cap = 60 if header == 'product_description' else 26
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), cap)

    ws.freeze_panes = 'A2'

    # Metadata sheet
    meta = wb.create_sheet('Metadata')
    meta.append(['Parameter', 'Value'])
    for cell in meta[1]:
        cell.font = header_font
        cell.fill = header_fill
    meta.append(['start_date', start_date])
    meta.append(['end_date (inclusive)', end_date])
    meta.append(['as_of_date (today)', as_of.isoformat()])
    meta.append(['SKU cohort',
                 'items received via voucher (vou_class=0, status=4, vou_type IN (0,1))'])
    meta.append(['row count', len(rows)])
    meta.append(['distinct UPCs', len({r['upc'] for r in rows if r['upc']})])
    meta.append(['distinct brands', len({r['brand_code'] for r in rows if r['brand_code']})])
    meta.append(['distinct seasons', len({r['season'] for r in rows if r['season']})])
    meta.column_dimensions['A'].width = 32
    meta.column_dimensions['B'].width = 74

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    parser.add_argument('--start-date', default='2025-01-01',
                        help='Inclusive lower bound (default: 2025-01-01)')
    parser.add_argument('--end-date', default=None,
                        help='Inclusive upper bound (default: today)')
    parser.add_argument('--out-dir', default='document/reports',
                        help='Output directory (default: document/reports)')
    args = parser.parse_args()

    as_of = date.today()
    end_date = args.end_date or as_of.isoformat()

    print(f'[1/2] Querying product master list ({args.start_date} -> {end_date}) ...')
    raw = fetch_rows(args.start_date, end_date)
    rows = shape_rows(raw)
    print(f'      -> {len(rows)} SKU rows')

    out_path = Path(args.out_dir) / f'product_master_list_{args.start_date}_{as_of.isoformat()}.xlsx'
    print(f'[2/2] Writing {out_path} ...')
    write_xlsx(rows, out_path, args.start_date, end_date, as_of)

    print(f'\nDone. Written to:')
    print(f'  {out_path.resolve()}')
    print(f'  SKU rows:         {len(rows):,}')
    print(f'  Distinct UPCs:    {len({r["upc"] for r in rows if r["upc"]}):,}')
    print(f'  Distinct brands:  {len({r["brand_code"] for r in rows if r["brand_code"]}):,}')
    print(f'  Distinct seasons: {len({r["season"] for r in rows if r["season"]}):,}')


if __name__ == '__main__':
    main()
