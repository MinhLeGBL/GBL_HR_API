"""
Master-data Excel report: monthly sales by store.

For each (year, month, store_code) in the window: net revenue before VAT
(sales minus returns, item_type 1 as +qty and item_type 2 as -qty), split
into FP vs MD by LINE discount only (di.DISC_PERC <= 30 => FP; > 30 => MD),
plus explicit sold/returned item counts.

VAT stripped by using actual per-line tax (di.price - di.tax_amt) — NOT the
flat 10% that commission uses (commission stays on price / 1.1 per CR #20;
this report intentionally diverges).

Filters: d.STATUS = 4, d.receipt_type IN (0, 1), di.ITEM_TYPE IN (1, 2).

Long layout: one row per (year_month, store_code).

Output: document/reports/monthly_sales_by_store_<start>_<today>.xlsx
"""
import argparse
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from app.core.database import get_oracle_connection


# Per-line before-VAT revenue with return-sign handling.
# Reused in FP, MD, and total revenue expressions.
_REV_LINE = (
    "(CASE WHEN di.ITEM_TYPE = 2 THEN di.qty * -1 ELSE di.qty END) "
    "* (di.price - di.tax_amt)"
)

QUERY = f"""
    SELECT
        EXTRACT(YEAR  FROM d.invc_post_date)                       AS year,
        EXTRACT(MONTH FROM d.invc_post_date)                       AS month,
        d.STORE_CODE                                               AS store_code,

        ROUND(SUM({_REV_LINE}), 0)                                 AS revenue_before_vat,

        ROUND(SUM(CASE WHEN di.DISC_PERC <= 30
                       THEN {_REV_LINE} ELSE 0 END), 0)            AS revenue_fp_before_vat,

        ROUND(SUM(CASE WHEN di.DISC_PERC > 30
                       THEN {_REV_LINE} ELSE 0 END), 0)            AS revenue_md_before_vat,

        ROUND(SUM(CASE WHEN di.ITEM_TYPE = 2
                       THEN di.qty * (di.price - di.tax_amt)
                       ELSE 0 END), 0)                             AS return_revenue_before_vat,

        SUM(CASE WHEN di.ITEM_TYPE = 1 AND di.DISC_PERC <= 30
                 THEN di.qty ELSE 0 END)                           AS items_sold_fp,

        SUM(CASE WHEN di.ITEM_TYPE = 1 AND di.DISC_PERC > 30
                 THEN di.qty ELSE 0 END)                           AS items_sold_md,

        SUM(CASE WHEN di.ITEM_TYPE = 1 THEN di.qty ELSE 0 END)     AS items_sold,

        SUM(CASE WHEN di.ITEM_TYPE = 2 THEN di.qty ELSE 0 END)     AS items_returned,

        COUNT(DISTINCT d.SID)                                      AS bill_count
    FROM DOCUMENT d
    JOIN DOCUMENT_ITEM di ON d.SID = di.DOC_SID
    WHERE d.STATUS = 4
      AND d.receipt_type IN (0, 1)
      AND di.ITEM_TYPE IN (1, 2)
      AND d.invc_post_date >= TO_DATE(:start_date, 'YYYY-MM-DD')
      AND TRUNC(d.invc_post_date) <= TO_DATE(:end_date, 'YYYY-MM-DD')
    GROUP BY EXTRACT(YEAR FROM d.invc_post_date),
             EXTRACT(MONTH FROM d.invc_post_date),
             d.STORE_CODE
    ORDER BY year, month, store_code
"""


def fetch_rows(start_date: str, end_date: str):
    conn = get_oracle_connection()
    if conn is None:
        raise RuntimeError('Could not open Oracle connection')
    try:
        cur = conn.cursor()
        try:
            cur.execute(QUERY, {'start_date': start_date, 'end_date': end_date})
            cols = [d[0].lower() for d in cur.description]
            rows = cur.fetchall()
        finally:
            cur.close()
    finally:
        conn.close()
    return [dict(zip(cols, r)) for r in rows]


def shape_rows(raw_rows):
    out = []
    for r in raw_rows:
        y = int(r['year'])
        m = int(r['month'])
        out.append({
            'year': y,
            'month': m,
            'year_month': f'{y:04d}-{m:02d}',
            'store_code': r.get('store_code') or '',
            'revenue_before_vat':        int(r.get('revenue_before_vat') or 0),
            'revenue_fp_before_vat':     int(r.get('revenue_fp_before_vat') or 0),
            'revenue_md_before_vat':     int(r.get('revenue_md_before_vat') or 0),
            'return_revenue_before_vat': int(r.get('return_revenue_before_vat') or 0),
            'items_sold_fp':             int(r.get('items_sold_fp') or 0),
            'items_sold_md':             int(r.get('items_sold_md') or 0),
            'items_sold':                int(r.get('items_sold') or 0),
            'items_returned':            int(r.get('items_returned') or 0),
            'bill_count':                int(r.get('bill_count') or 0),
        })
    return out


def write_xlsx(rows, out_path: Path, start_date: str, end_date: str, as_of: date):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Monthly Sales by Store'

    headers = [
        'year', 'month', 'year_month', 'store_code',
        'revenue_before_vat',
        'revenue_fp_before_vat', 'revenue_md_before_vat',
        'return_revenue_before_vat',
        'items_sold_fp', 'items_sold_md', 'items_sold', 'items_returned',
        'bill_count',
    ]
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='305496', end_color='305496', fill_type='solid')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for r in rows:
        ws.append([r[h] for h in headers])

    # Format revenue columns as #,##0
    rev_cols = [i + 1 for i, h in enumerate(headers) if 'revenue' in h]
    for row_idx in range(2, ws.max_row + 1):
        for c in rev_cols:
            ws.cell(row=row_idx, column=c).number_format = '#,##0'

    # Autosize
    for col_idx, header in enumerate(headers, start=1):
        cells = [ws.cell(row=r, column=col_idx).value
                 for r in range(2, ws.max_row + 1)] if ws.max_row > 1 else []
        max_len = max(len(str(header)),
                      *(len(f'{v:,}') if isinstance(v, int) else len(str(v or ''))
                        for v in cells)) if cells else len(header)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 26)

    ws.freeze_panes = 'A2'

    # Metadata sheet
    meta = wb.create_sheet('Metadata')
    meta.append(['Parameter', 'Value'])
    for cell in meta[1]:
        cell.font = header_font
        cell.fill = header_fill
    meta.append(['start_date', start_date])
    meta.append(['end_date (inclusive)', end_date])
    meta.append(['as_of_date (today)', as_of.isoformat()])
    meta.append(['revenue basis', 'before VAT — di.price - di.tax_amt (actual per-line VAT)'])
    meta.append(['FP/MD rule', 'line discount only — di.DISC_PERC <= 30 = FP'])
    meta.append(['returns handled', 'net of returns (item_type=2 negates qty)'])
    meta.append(['filters', 'd.STATUS=4, d.receipt_type IN (0,1), di.ITEM_TYPE IN (1,2)'])
    meta.append(['row count', len(rows)])
    meta.append(['total revenue (before VAT)',
                 sum(r['revenue_before_vat'] for r in rows)])
    meta.append(['total FP revenue',
                 sum(r['revenue_fp_before_vat'] for r in rows)])
    meta.append(['total MD revenue',
                 sum(r['revenue_md_before_vat'] for r in rows)])
    meta.append(['total return revenue',
                 sum(r['return_revenue_before_vat'] for r in rows)])
    meta.append(['distinct months',
                 len({r['year_month'] for r in rows})])
    meta.append(['distinct stores',
                 len({r['store_code'] for r in rows})])
    meta.column_dimensions['A'].width = 34
    meta.column_dimensions['B'].width = 62

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    parser.add_argument('--start-date', default='2025-01-01',
                        help='Inclusive lower bound on invc_post_date (default: 2025-01-01)')
    parser.add_argument('--end-date', default=None,
                        help='Inclusive upper bound on invc_post_date (default: today)')
    parser.add_argument('--out-dir', default='document/reports',
                        help='Output directory (default: document/reports)')
    args = parser.parse_args()

    as_of = date.today()
    end_date = args.end_date or as_of.isoformat()

    print(f'[1/2] Querying monthly sales by store ({args.start_date} -> {end_date}) ...')
    raw = fetch_rows(args.start_date, end_date)
    rows = shape_rows(raw)
    print(f'      -> {len(rows)} (year_month, store) rows')

    out_path = Path(args.out_dir) / f'monthly_sales_by_store_{args.start_date}_{as_of.isoformat()}.xlsx'
    print(f'[2/2] Writing {out_path} ...')
    write_xlsx(rows, out_path, args.start_date, end_date, as_of)

    total_rev = sum(r['revenue_before_vat'] for r in rows)
    total_fp = sum(r['revenue_fp_before_vat'] for r in rows)
    total_md = sum(r['revenue_md_before_vat'] for r in rows)
    print(f'\nDone. Written to:')
    print(f'  {out_path.resolve()}')
    print(f'  Rows:             {len(rows)}')
    print(f'  Total revenue:    {total_rev:,}')
    print(f'    of which FP:    {total_fp:,}')
    print(f'    of which MD:    {total_md:,}')
    print(f'  Invariant check:  FP + MD = total?  {total_fp + total_md == total_rev}')
    print(f'  Distinct months:  {len({r["year_month"] for r in rows})}')
    print(f'  Distinct stores:  {len({r["store_code"] for r in rows})}')


if __name__ == '__main__':
    main()
