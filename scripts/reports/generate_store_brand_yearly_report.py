"""
Master-data Excel report: yearly revenue / transactions / units by store x brand,
with an origin split on the units sold.

Grain: one row per (year, store, brand).

Metrics
  revenue_before_vat   net of returns, bill discount applied:
                         signed_qty * (di.price - di.tax_amt) * (1 - d.DISC_PERC/100)
                       di.price is already net of the LINE discount; only the
                       bill-level discount has to be applied on top.
  transactions         distinct SALE receipts (receipt_type = 0) containing at
                       least one line of that brand at that store
  items_sold           units on item_type = 1 lines
  items_returned       units on item_type = 2 lines
  items_net            items_sold - items_returned

Origin split (on items_sold only — returns are excluded, they muddy origin)
  How stock reaches a store, as Retail Pro records it here:
    - vendor receipt = VOUCHER with VOU_CLASS=0, VOU_TYPE=0, STATUS=4 and
      SLIP_FLAG=0. 815 of 819 such vouchers post to HQ, so HQ is effectively
      the single import point.
    - a transfer-in ALSO writes a VOUCHER, but with SLIP_FLAG=1. Those are
      NOT receipts and are deliberately ignored here — the movement itself is
      read from SLIP / SLIP_ITEM instead.

  items_original       the store is where the unit first landed: either the
                       store took it on a vendor receipt directly, or the
                       FIRST slip that moved this item into this store came
                       from HQ. This is "assigned to the store at the
                       beginning" — one transfer since import.
  items_transferred    the first slip into this store came from another RETAIL
                       store, i.e. the unit was moved on after its first
                       posting (two or more transfers since import)
  items_unknown        no vendor receipt and no inbound slip for this
                       (item, store). Concentrated in 2024 — legacy stock
                       loaded at the Dec-2023 go-live via inventory
                       adjustments, which carry no origin store.
  items_mixed_flagged  subset of items_sold whose (item, store) pair received
                       the SKU BOTH from HQ and from another store. First
                       inbound decides the bucket, but the split is a guess for
                       these — reported so the exposure is visible.

Origin is resolved at (SKU, store) grain, NOT per physical unit: nothing in
this database is serialised (0 of 137,259 receipt lines, 0 of 105,844 slip
lines and 0 of 82,734 sale lines carry a SERIAL_NO), so individual units
cannot be traced. See items_mixed_flagged for how often that bites.

Filters: d.STATUS = 4, d.receipt_type IN (0, 1), di.ITEM_TYPE IN (1, 2).

Output: document/reports/store_brand_yearly_<start>_<end>_<today>.xlsx
"""
import argparse
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from app.core.database import get_oracle_connection


# Per-line net revenue before VAT, returns signed negative, bill discount applied.
_REV_LINE = (
    "(CASE WHEN di.ITEM_TYPE = 2 THEN di.qty * -1 ELSE di.qty END) "
    "* (di.price - NVL(di.tax_amt, 0)) "
    "* (1 - NVL(d.DISC_PERC, 0) / 100)"
)

# Units on sale lines only — origin buckets are counted against these.
_SOLD_QTY = "CASE WHEN di.ITEM_TYPE = 1 THEN di.qty ELSE 0 END"

QUERY = f"""
WITH first_inbound AS (
    -- Earliest slip that moved each item into each store.
    SELECT item_sid, in_store_sid, out_store_sid
    FROM (
        SELECT
            si.ITEM_SID          AS item_sid,
            sl.IN_STORE_SID      AS in_store_sid,
            sl.OUT_STORE_SID     AS out_store_sid,
            ROW_NUMBER() OVER (PARTITION BY si.ITEM_SID, sl.IN_STORE_SID
                               ORDER BY sl.POST_DATE, sl.SID) AS rn
        FROM SLIP sl
        JOIN SLIP_ITEM si ON si.SLIP_SID = sl.SID
        WHERE sl.STATUS = 4
          AND sl.IN_STORE_SID IS NOT NULL
    )
    WHERE rn = 1
),
inbound_mix AS (
    -- Did this (item, store) ever receive the SKU from HQ AND from a store?
    SELECT
        si.ITEM_SID      AS item_sid,
        sl.IN_STORE_SID  AS in_store_sid,
        MAX(CASE WHEN so.STORE_CODE = :hq_code THEN 1 ELSE 0 END) AS from_hq,
        MAX(CASE WHEN so.STORE_CODE <> :hq_code THEN 1 ELSE 0 END) AS from_store
    FROM SLIP sl
    JOIN SLIP_ITEM si  ON si.SLIP_SID = sl.SID
    LEFT JOIN STORE so ON so.SID = sl.OUT_STORE_SID
    WHERE sl.STATUS = 4
      AND sl.IN_STORE_SID IS NOT NULL
    GROUP BY si.ITEM_SID, sl.IN_STORE_SID
),
vendor_receipt AS (
    -- True vendor receipts only. SLIP_FLAG = 1 vouchers are transfer-ins.
    SELECT DISTINCT vi.ITEM_SID AS item_sid, v.STORE_SID AS store_sid
    FROM VOUCHER v
    JOIN VOU_ITEM vi ON vi.VOU_SID = v.SID
    WHERE v.VOU_CLASS = 0
      AND v.VOU_TYPE = 0
      AND v.STATUS = 4
      AND NVL(v.SLIP_FLAG, 0) = 0
),
sale_lines AS (
    SELECT
        EXTRACT(YEAR FROM d.invc_post_date)        AS yr,
        d.STORE_CODE                               AS store_code,
        NVL(vb.VEND_NAME, di.VEND_CODE)            AS brand,
        d.SID                                      AS doc_sid,
        d.receipt_type                             AS receipt_type,
        di.ITEM_TYPE                               AS item_type,
        di.qty                                     AS qty,
        {_REV_LINE}                                AS rev,
        {_SOLD_QTY}                                AS sold_qty,
        CASE
            WHEN vr.item_sid IS NOT NULL                 THEN 'ORIGINAL'
            WHEN fi.out_store_sid IS NULL                THEN 'UNKNOWN'
            WHEN so.STORE_CODE = :hq_code                THEN 'ORIGINAL'
            ELSE 'TRANSFERRED'
        END                                        AS origin,
        CASE WHEN im.from_hq = 1 AND im.from_store = 1
             THEN 1 ELSE 0 END                     AS mixed_flag
    FROM DOCUMENT d
    JOIN DOCUMENT_ITEM di        ON di.DOC_SID = d.SID
    JOIN STORE st                ON st.STORE_CODE = d.STORE_CODE
    LEFT JOIN INVN_SBS_ITEM i    ON i.SID = di.INVN_SBS_ITEM_SID
    LEFT JOIN VENDOR vb          ON vb.SID = i.VEND_SID
    LEFT JOIN vendor_receipt vr  ON vr.item_sid = di.INVN_SBS_ITEM_SID
                                AND vr.store_sid = st.SID
    LEFT JOIN first_inbound fi   ON fi.item_sid = di.INVN_SBS_ITEM_SID
                                AND fi.in_store_sid = st.SID
    LEFT JOIN STORE so           ON so.SID = fi.out_store_sid
    LEFT JOIN inbound_mix im     ON im.item_sid = di.INVN_SBS_ITEM_SID
                                AND im.in_store_sid = st.SID
    WHERE d.STATUS = 4
      AND d.receipt_type IN (0, 1)
      AND di.ITEM_TYPE IN (1, 2)
      AND d.invc_post_date >= TO_DATE(:start_date, 'YYYY-MM-DD')
      AND TRUNC(d.invc_post_date) <= TO_DATE(:end_date, 'YYYY-MM-DD')
)
SELECT
    sl.yr                                                        AS year,
    sl.store_code                                                AS store_code,
    NVL(st.STORE_NAME, sl.store_code)                            AS store_name,
    NVL(sl.brand, '(Unknown)')                                   AS brand,

    ROUND(SUM(sl.rev), 0)                                        AS revenue_before_vat,
    COUNT(DISTINCT CASE WHEN sl.receipt_type = 0
                        THEN sl.doc_sid END)                     AS transactions,

    SUM(sl.sold_qty)                                             AS items_sold,
    SUM(CASE WHEN sl.item_type = 2 THEN sl.qty ELSE 0 END)       AS items_returned,
    SUM(CASE WHEN sl.item_type = 1 THEN sl.qty ELSE -sl.qty END) AS items_net,

    SUM(CASE WHEN sl.origin = 'ORIGINAL'
             THEN sl.sold_qty ELSE 0 END)                        AS items_original,
    SUM(CASE WHEN sl.origin = 'TRANSFERRED'
             THEN sl.sold_qty ELSE 0 END)                        AS items_transferred,
    SUM(CASE WHEN sl.origin = 'UNKNOWN'
             THEN sl.sold_qty ELSE 0 END)                        AS items_unknown,
    SUM(CASE WHEN sl.mixed_flag = 1
             THEN sl.sold_qty ELSE 0 END)                        AS items_mixed_flagged
FROM sale_lines sl
LEFT JOIN STORE st ON st.STORE_CODE = sl.store_code
GROUP BY sl.yr, sl.store_code, NVL(st.STORE_NAME, sl.store_code),
         NVL(sl.brand, '(Unknown)')
ORDER BY sl.yr, sl.store_code, revenue_before_vat DESC
"""


# Transactions are counted per (year, store, brand), so a receipt spanning two
# brands is counted once under each. That is correct per brand but must NOT be
# summed for a store total, so the rollup takes its count from here instead.
QUERY_STORE_TXN = """
    SELECT
        EXTRACT(YEAR FROM d.invc_post_date)  AS year,
        d.STORE_CODE                         AS store_code,
        COUNT(DISTINCT CASE WHEN d.receipt_type = 0
                            THEN d.SID END)  AS transactions
    FROM DOCUMENT d
    JOIN DOCUMENT_ITEM di ON di.DOC_SID = d.SID
    WHERE d.STATUS = 4
      AND d.receipt_type IN (0, 1)
      AND di.ITEM_TYPE IN (1, 2)
      AND d.invc_post_date >= TO_DATE(:start_date, 'YYYY-MM-DD')
      AND TRUNC(d.invc_post_date) <= TO_DATE(:end_date, 'YYYY-MM-DD')
    GROUP BY EXTRACT(YEAR FROM d.invc_post_date), d.STORE_CODE
"""

def _fetch(query, params):
    conn = get_oracle_connection()
    if conn is None:
        raise RuntimeError('Could not open Oracle connection')
    try:
        cur = conn.cursor()
        try:
            cur.execute(query, params)
            cols = [d[0].lower() for d in cur.description]
            rows = cur.fetchall()
        finally:
            cur.close()
    finally:
        conn.close()
    return [dict(zip(cols, r)) for r in rows]


def fetch_rows(start_date: str, end_date: str, hq_code: str):
    return _fetch(QUERY, {'start_date': start_date,
                          'end_date': end_date,
                          'hq_code': hq_code})


def fetch_store_txn(start_date: str, end_date: str):
    """True distinct sale-receipt count per (year, store)."""
    rows = _fetch(QUERY_STORE_TXN, {'start_date': start_date, 'end_date': end_date})
    return {(r['year'], r['store_code']): r['transactions'] for r in rows}


COLUMNS = [
    ('year', 'year', '0'),
    ('store_code', 'store_code', None),
    ('store_name', 'store_name', None),
    ('brand', 'brand', None),
    ('revenue_before_vat', 'revenue_before_vat', '#,##0'),
    ('transactions', 'transactions', '#,##0'),
    ('items_sold', 'items_sold', '#,##0'),
    ('items_returned', 'items_returned', '#,##0'),
    ('items_net', 'items_net', '#,##0'),
    ('items_original', 'items_original', '#,##0'),
    ('items_transferred', 'items_transferred', '#,##0'),
    ('items_unknown', 'items_unknown', '#,##0'),
    ('pct_original', None, '0.0"%"'),
    ('items_mixed_flagged', 'items_mixed_flagged', '#,##0'),
]


def _pct_original(r):
    sold = r['items_sold'] or 0
    return round((r['items_original'] or 0) / sold * 100, 1) if sold else 0.0


def summarise_by_store(rows, store_txn):
    """Roll the (year, store, brand) rows up to (year, store)."""
    acc = {}
    for r in rows:
        key = (r['year'], r['store_code'], r['store_name'])
        a = acc.setdefault(key, {k: 0 for _, k, _ in COLUMNS if k and k not in
                                ('year', 'store_code', 'store_name', 'brand')})
        a.setdefault('brands', set()).add(r['brand'])
        for k in ('revenue_before_vat', 'transactions', 'items_sold',
                  'items_returned', 'items_net', 'items_original',
                  'items_transferred', 'items_unknown', 'items_mixed_flagged'):
            a[k] += r[k] or 0
    out = []
    for (yr, code, name), a in sorted(acc.items()):
        brands = a.pop('brands')
        row = {'year': yr, 'store_code': code, 'store_name': name,
               'brand_count': len(brands)}
        row.update(a)
        # Overwrite the summed-across-brands figure with a true distinct count.
        row['transactions'] = store_txn.get((yr, code), 0)
        out.append(row)
    return out


def write_xlsx(rows, store_rows, out_path: Path, start_date, end_date, as_of, hq_code):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='305496', end_color='305496', fill_type='solid')

    def _style_header(ws):
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

    wb = Workbook()

    # ---- Sheet 1: detail (year, store, brand) ----
    ws = wb.active
    ws.title = 'By Store & Brand'
    ws.append([label for label, _, _ in COLUMNS])
    _style_header(ws)
    for r in rows:
        ws.append([_pct_original(r) if key is None else r[key]
                   for _, key, _ in COLUMNS])
    for col_idx, (_, _, fmt) in enumerate(COLUMNS, start=1):
        if fmt:
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row=row_idx, column=col_idx).number_format = fmt
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = ws.dimensions
    for col_idx, (label, _, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(len(label) + 3, 12)

    # ---- Sheet 2: store rollup ----
    store_cols = ['year', 'store_code', 'store_name', 'brand_count',
                  'revenue_before_vat', 'transactions', 'items_sold',
                  'items_returned', 'items_net', 'items_original',
                  'items_transferred', 'items_unknown', 'pct_original',
                  'items_mixed_flagged']
    ws2 = wb.create_sheet('By Store')
    ws2.append(store_cols)
    _style_header(ws2)
    for r in store_rows:
        ws2.append([_pct_original(r) if c == 'pct_original' else r[c]
                    for c in store_cols])
    for col_idx, c in enumerate(store_cols, start=1):
        fmt = ('0.0"%"' if c == 'pct_original'
               else '#,##0' if c not in ('year', 'store_code', 'store_name') else None)
        if fmt:
            for row_idx in range(2, ws2.max_row + 1):
                ws2.cell(row=row_idx, column=col_idx).number_format = fmt
        ws2.column_dimensions[get_column_letter(col_idx)].width = max(len(c) + 3, 12)
    ws2.freeze_panes = 'D2'

    # ---- Sheet 3: metadata ----
    meta = wb.create_sheet('Metadata')
    meta.append(['Parameter', 'Value'])
    for cell in meta[1]:
        cell.font = header_font
        cell.fill = header_fill
    entries = [
        ('window', f'{start_date} .. {end_date}'),
        ('as_of_date', as_of.isoformat()),
        ('grain', 'one row per (year, store, brand) on sheet 1'),
        ('brand', 'VENDOR.VEND_NAME via INVN_SBS_ITEM.VEND_SID (100% join coverage)'),
        ('revenue', 'before VAT, returns netted: signed_qty * (di.price - di.tax_amt) '
                    '* (1 - d.DISC_PERC/100)'),
        ('transactions', 'distinct SALE receipts (receipt_type = 0) containing >= 1 line '
                         'of that brand; returns excluded from the count. A receipt '
                         'spanning two brands counts once under EACH brand, so the '
                         'per-brand column does not sum to the store total — the By Store '
                         'sheet carries a true distinct count instead'),
        ('items_sold', 'units on item_type = 1 lines (returns reported separately)'),
        ('import point', f'{hq_code} — 815 of 819 vendor receipts (VOUCHER SLIP_FLAG=0) '
                         f'post to {hq_code}; the other 4 are RWT/RWD direct receipts'),
        ('transfer-in vouchers', 'VOUCHER rows with SLIP_FLAG=1 are transfer receipts, '
                                 'NOT vendor receipts — excluded from origin logic; '
                                 'movement is read from SLIP/SLIP_ITEM instead'),
        ('items_original', f'store took the SKU on a vendor receipt directly, OR the first '
                           f'slip into this store came from {hq_code} — i.e. one transfer '
                           f'since import ("assigned to the store at the beginning")'),
        ('items_transferred', 'first slip into this store came from another retail store — '
                              'two or more transfers since import'),
        ('items_unknown', 'no vendor receipt and no inbound slip for this (item, store). '
                          'Concentrated in 2024 — legacy stock loaded at the Dec-2023 '
                          'go-live via inventory adjustments, which carry no origin store'),
        ('items_mixed_flagged', 'units whose (item, store) pair received the SKU from BOTH '
                                f'{hq_code} and another store (8.5% of all pairs). First '
                                'inbound decides the bucket; for these it is an estimate'),
        ('serialisation caveat', 'nothing in this database is serialised (0 SERIAL_NO on '
                                 '137,259 receipt lines / 105,844 slip lines / 82,734 sale '
                                 'lines), so origin is resolved at (SKU, store) grain, not '
                                 'per physical unit'),
        ('filters', 'd.STATUS=4, d.receipt_type IN (0,1), di.ITEM_TYPE IN (1,2); '
                    'SLIP.STATUS=4; VOUCHER VOU_CLASS=0, VOU_TYPE=0, STATUS=4'),
    ]
    for k, v in entries:
        meta.append([k, v])
    meta.column_dimensions['A'].width = 24
    meta.column_dimensions['B'].width = 100
    for row in meta.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical='top')

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    parser.add_argument('--start-date', default='2024-01-01')
    parser.add_argument('--end-date', default='2026-12-31')
    parser.add_argument('--hq-code', default='HQ',
                        help='Store code of the import point (default: HQ)')
    parser.add_argument('--out-dir', default='document/reports')
    args = parser.parse_args()

    as_of = date.today()

    print(f'[1/2] Querying store x brand aggregates ({args.start_date} .. {args.end_date}) ...')
    rows = fetch_rows(args.start_date, args.end_date, args.hq_code)
    print(f'      -> {len(rows)} (year, store, brand) rows')
    store_txn = fetch_store_txn(args.start_date, args.end_date)
    store_rows = summarise_by_store(rows, store_txn)

    out_path = (Path(args.out_dir) /
                f'store_brand_yearly_{args.start_date}_{args.end_date}_{as_of.isoformat()}.xlsx')
    print(f'[2/2] Writing {out_path} ...')
    write_xlsx(rows, store_rows, out_path, args.start_date, args.end_date,
               as_of, args.hq_code)

    print('\nDone. Written to:')
    print(f'  {out_path.resolve()}\n')
    for r in store_rows:
        sold = r['items_sold'] or 1
        print(f"  {r['year']} {r['store_code']:<4} rev={r['revenue_before_vat']:>16,.0f} "
              f"txn={r['transactions']:>6,} sold={r['items_sold']:>6,} "
              f"orig={r['items_original']:>6,} ({r['items_original']/sold*100:5.1f}%) "
              f"xfer={r['items_transferred']:>6,} unk={r['items_unknown']:>6,}")


if __name__ == '__main__':
    main()
