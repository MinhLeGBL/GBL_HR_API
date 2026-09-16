"""
Excel report: average on-hand inventory price vs realized selling price, by month,
for one or more stores — H1 2026. One sheet per store.

Per period (month + H1 total), each sheet shows:
  avg_onhand_qty          average of DAILY on-hand units over the period
                          (reconstructed — Retail Pro stores no daily history)
  onhand_avg_list_price   unit-days-weighted average RETAIL (list) price of that
                          on-hand inventory (INVN_SBS_PRICE level 1 'Retail')
  units_sold              units sold in the period (ITEM_TYPE=1)
  avg_selling_price_sold  qty-weighted ACTUAL selling price of those units
                          (di.PRICE — after the sale discount)

On-hand reconstruction: on_hand(day) walked forward from a Jan-1 opening derived
from the current snapshot minus in-window movements. Movements affecting a
store's on-hand:
  - sales / customer returns   DOCUMENT.STORE_SID   (item_type 1 −, 2 +)
  - receipts / vendor returns  VOUCHER.STORE_SID    (vou_class 0, type 0 +, 1 −)
  - transfers OUT of the store VOUCHER.ORIG_STORE_SID = store, STORE_SID ≠ store
                               (recorded as a receiving voucher at the
                                destination; −qty here)
No adjustments / vou_class=2 transfers exist for these stores in the window.

All prices VAT-inclusive. Output: document/reports/store_onhand_avg_price_H1_2026_<today>.xlsx
"""
import argparse
import sys
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from app.core.database import get_oracle_connection

START, END_EXCL, END = date(2026, 1, 1), date(2026, 7, 1), date(2026, 6, 30)
MONTHS = ['2026-01', '2026-02', '2026-03', '2026-04', '2026-05', '2026-06']
DEFAULT_STORES = ['RWD', 'RWT']


def _dt(v):
    return v.date() if hasattr(v, 'date') else v


def resolve_stores(cur, codes):
    cur.execute(
        "SELECT STORE_CODE, SID, STORE_NAME FROM RPS.STORE WHERE STORE_CODE IN (%s)"
        % ','.join(f"'{c}'" for c in codes))
    return {code: (sid, name) for code, sid, name in cur.fetchall()}


def fetch_onhand_inputs(cur, sid):
    cur.execute("""
    WITH price1 AS (SELECT p.INVN_SBS_ITEM_SID item_sid, p.PRICE
      FROM RPS.INVN_SBS_PRICE p JOIN RPS.PRICE_LEVEL pl ON pl.SID=p.PRICE_LVL_SID
      WHERE pl.PRICE_LVL=1)
    SELECT q.INVN_SBS_ITEM_SID, q.QTY, pr.PRICE
    FROM RPS.INVN_SBS_ITEM_QTY q LEFT JOIN price1 pr ON pr.item_sid=q.INVN_SBS_ITEM_SID
    WHERE q.STORE_SID=:s""", {'s': sid})
    cur_qty, price = defaultdict(int), {}
    for item, qty, pr in cur.fetchall():
        cur_qty[int(item)] += int(qty or 0)
        if pr is not None:
            price[int(item)] = float(pr)

    moves = []
    cur.execute("""SELECT b.INVN_SBS_ITEM_SID, TRUNC(a.invc_post_date),
           CASE WHEN b.item_type=2 THEN b.qty ELSE -b.qty END
        FROM RPS.DOCUMENT a JOIN RPS.DOCUMENT_ITEM b ON b.DOC_SID=a.SID
        WHERE a.STORE_SID=:s AND a.status=4 AND a.receipt_type IN (0,1)
          AND b.item_type IN (1,2) AND a.invc_post_date>=DATE'2026-01-01'""", {'s': sid})
    moves += [(int(i), _dt(d), int(q)) for i, d, q in cur.fetchall()]
    cur.execute("""SELECT vi.item_sid, TRUNC(vh.post_date),
           CASE WHEN vh.vou_type=1 THEN -vi.qty ELSE vi.qty END
        FROM RPS.VOUCHER vh JOIN RPS.VOU_ITEM vi ON vi.vou_sid=vh.SID
        WHERE vh.STORE_SID=:s AND vh.status=4 AND vh.vou_class=0
          AND vh.vou_type IN (0,1) AND vh.post_date>=DATE'2026-01-01'""", {'s': sid})
    moves += [(int(i), _dt(d), int(q)) for i, d, q in cur.fetchall()]
    cur.execute("""SELECT vi.item_sid, TRUNC(vh.post_date), -vi.qty
        FROM RPS.VOUCHER vh JOIN RPS.VOU_ITEM vi ON vi.vou_sid=vh.SID
        WHERE vh.ORIG_STORE_SID=:s AND vh.STORE_SID<>:s AND vh.status=4
          AND vh.vou_class=0 AND vh.post_date>=DATE'2026-01-01'""", {'s': sid})
    moves += [(int(i), _dt(d), int(q)) for i, d, q in cur.fetchall()]
    return cur_qty, price, moves


def monthly_onhand(cur_qty, price, moves):
    """Return {ym: (avg_daily_units, unit_days_wt_price)}, ('H1', ...), worst_neg.

    On-hand is floored at 0 per item when aggregating: a negative reconstructed
    on-hand means a within-window timing artifact (an outflow dated before its
    receipt lands), and physical stock can't be negative. The internal running
    value stays un-clamped so a later receipt correctly restores the item.
    """
    sum_moves = defaultdict(int)
    for it, d, delta in moves:
        sum_moves[it] += delta
    onhand = {it: cur_qty.get(it, 0) - sum_moves[it]      # opening at end 2025-12-31
              for it in set(cur_qty) | set(sum_moves)}

    by_day = defaultdict(list)
    for it, d, delta in moves:
        by_day[d].append((it, delta))
    active = [it for it in onhand if it in price]          # priced items only

    m_units, m_value, m_days = defaultdict(float), defaultdict(float), defaultdict(int)
    worst_neg = 0
    d = START
    while d <= END:
        for it, delta in by_day.get(d, []):
            onhand[it] = onhand.get(it, 0) + delta
        du = dv = 0.0
        neg = 0
        for it in active:
            q = onhand[it]
            if q > 0:
                du += q
                dv += price[it] * q
            elif q < 0:
                neg -= q
        worst_neg = max(worst_neg, neg)
        ym = d.strftime('%Y-%m')
        m_units[ym] += du
        m_value[ym] += dv
        m_days[ym] += 1
        d += timedelta(days=1)

    out = {}
    for ym in MONTHS:
        u, v, n = m_units[ym], m_value[ym], m_days[ym]
        out[ym] = (u / n if n else 0, v / u if u else 0)
    tu, tv, td = sum(m_units.values()), sum(m_value.values()), sum(m_days.values())
    out['H1'] = (tu / td if td else 0, tv / tu if tu else 0)
    return out, worst_neg


def fetch_sold(cur, sid):
    """Return {ym|'H1': (units, avg_selling_price_after_discount)}."""
    cur.execute("""
        SELECT NVL(TO_CHAR(a.invc_post_date,'YYYY-MM'),'H1') ym,
               SUM(b.qty) units,
               ROUND(SUM(b.PRICE*b.qty)/SUM(b.qty)) avg_sold
        FROM RPS.DOCUMENT a JOIN RPS.DOCUMENT_ITEM b ON b.DOC_SID=a.SID
        WHERE a.STORE_SID=:s AND a.status=4 AND a.receipt_type IN (0,1) AND b.item_type=1
          AND a.invc_post_date>=DATE'2026-01-01' AND a.invc_post_date<DATE'2026-07-01'
        GROUP BY ROLLUP(TO_CHAR(a.invc_post_date,'YYYY-MM'))""", {'s': sid})
    return {ym: (int(u or 0), int(av or 0)) for ym, u, av in cur.fetchall()}


def build_rows(onhand, sold):
    rows = []
    for period in MONTHS + ['H1']:
        au, ap = onhand.get(period, (0, 0))
        su, sp = sold.get(period, (0, 0))
        rows.append({
            'period': 'H1 TOTAL' if period == 'H1' else period,
            'avg_onhand_qty': round(au),
            'onhand_avg_list_price': round(ap),
            'units_sold': su,
            'avg_selling_price_sold': sp,
        })
    return rows


HEADERS = [('period', 'Period'), ('avg_onhand_qty', 'Avg on-hand qty'),
           ('onhand_avg_list_price', 'On-hand avg list price'),
           ('units_sold', 'Units sold'),
           ('avg_selling_price_sold', 'Avg selling price (sold)')]


def write_xlsx(store_rows, out_path, as_of):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    hfont = Font(bold=True, color='FFFFFF')
    hfill = PatternFill(start_color='305496', end_color='305496', fill_type='solid')
    bold = Font(bold=True)
    title_font = Font(bold=True, size=13)

    wb = Workbook()
    first = True
    for code, (name, rows) in store_rows.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = code
        first = False

        ws.append([f'{code} — {name}: average on-hand price vs realized selling price, H1 2026'])
        ws['A1'].font = title_font
        ws.append([])
        ws.append([h[1] for h in HEADERS])
        for cell in ws[3]:
            cell.font = hfont; cell.fill = hfill
            cell.alignment = Alignment(horizontal='center')

        for r in rows:
            ws.append([r[k] for k, _ in HEADERS])
            if r['period'] == 'H1 TOTAL':
                for cell in ws[ws.max_row]:
                    cell.font = bold

        money = [i + 1 for i, (k, _) in enumerate(HEADERS)
                 if k in ('onhand_avg_list_price', 'avg_selling_price_sold')]
        qty = [i + 1 for i, (k, _) in enumerate(HEADERS)
               if k in ('avg_onhand_qty', 'units_sold')]
        for ri in range(4, ws.max_row + 1):
            for c in money:
                ws.cell(row=ri, column=c).number_format = '#,##0'
            for c in qty:
                ws.cell(row=ri, column=c).number_format = '#,##0'

        widths = [12, 16, 24, 12, 26]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A4'

        note = ('On-hand = avg of daily on-hand over the period (reconstructed), priced at unit-days-weighted '
                'RETAIL list price. Selling price = actual di.PRICE after discount. VAT-inclusive.')
        ws.append([]); ws.append([note])
        ws.cell(row=ws.max_row, column=1).font = Font(italic=True, size=9)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser(description='Store on-hand avg price vs selling price, H1 2026')
    ap.add_argument('--stores', nargs='+', default=DEFAULT_STORES)
    ap.add_argument('--out-dir', default='document/reports')
    args = ap.parse_args()
    as_of = date.today()

    conn = get_oracle_connection()
    if conn is None:
        raise RuntimeError('Could not open Oracle connection')
    store_rows = {}
    try:
        cur = conn.cursor()
        try:
            resolved = resolve_stores(cur, args.stores)
            for code in args.stores:
                if code not in resolved:
                    print(f'  ! store {code} not found — skipping'); continue
                sid, name = resolved[code]
                print(f'[{code}] {name} — reconstructing on-hand + sold ...')
                cur_qty, price, moves = fetch_onhand_inputs(cur, sid)
                onhand, worst_neg = monthly_onhand(cur_qty, price, moves)
                sold = fetch_sold(cur, sid)
                rows = build_rows(onhand, sold)
                store_rows[code] = (name, rows)
                h1 = rows[-1]
                print(f"      H1 avg on-hand {h1['avg_onhand_qty']:,} @ {h1['onhand_avg_list_price']:,} list | "
                      f"sold {h1['units_sold']:,} @ {h1['avg_selling_price_sold']:,} | recon worst-neg {worst_neg} units")
        finally:
            cur.close()
    finally:
        conn.close()

    out_path = Path(args.out_dir) / f'store_onhand_avg_price_H1_2026_{as_of.isoformat()}.xlsx'
    write_xlsx(store_rows, out_path, as_of)
    print(f'\nWrote {out_path.resolve()}  ({len(store_rows)} sheets)')


if __name__ == '__main__':
    main()
