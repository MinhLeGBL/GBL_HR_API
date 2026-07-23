"""
One-off Excel report: dormant customers — two sheets.

Sheet 1 — Lapsed
  Customers whose LAST purchase was strictly before --cutoff-date
  (default: 2025-06-15). One row per customer:
    name, phone, last_sale_date, last_store

Sheet 2 — Never Purchased
  Customers in the system (CUSTOMER table) with NO sales record.
  Same column shape; last_sale_date and last_store are blank.

Exclusions on both sheets: SYSADMIN, Tourist, customers with NULL name.

Output: document/reports/dormant_customers_<cutoff>_<today>.xlsx
"""
import argparse
import sys
from datetime import date
from pathlib import Path

# Allow running this file directly from the repo root.
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from app.core.database import get_oracle_connection


# Phone lives on rps.CUSTOMER_PHONE (one customer can have multiple). We pick
# the PRIMARY_FLAG=1 row first, breaking ties by lowest SEQ_NO. Same CTE is
# reused in both sheet queries via inline copy.
_PRIMARY_PHONE_CTE = """
    primary_phone AS (
        SELECT cust_sid, phone_no,
               ROW_NUMBER() OVER (
                   PARTITION BY cust_sid
                   ORDER BY PRIMARY_FLAG DESC, SEQ_NO ASC
               ) AS rn
        FROM CUSTOMER_PHONE
        WHERE phone_no IS NOT NULL
    )
"""

# Sheet 1 — customers with sales, but last sale strictly before :cutoff_date.
# ROW_NUMBER over invc_post_date DESC picks the latest bill per customer; tied
# bills broken by DOCUMENT.SID DESC for determinism. The outer WHERE then keeps
# only customers whose latest sale is before the cutoff.
QUERY_LAPSED = f"""
    WITH ranked_sales AS (
        SELECT
            d.BT_CUID         AS customer_sid,
            d.invc_post_date  AS sale_date,
            d.STORE_CODE      AS store_code,
            ROW_NUMBER() OVER (
                PARTITION BY d.BT_CUID
                ORDER BY d.invc_post_date DESC, d.SID DESC
            ) AS rn
        FROM DOCUMENT d
        WHERE d.STATUS = 4
          AND d.receipt_type IN (0, 1)
          AND d.BT_CUID IS NOT NULL
    ),
    {_PRIMARY_PHONE_CTE}
    SELECT
        TRIM(c.FIRST_NAME)         AS name,
        pp.phone_no                AS phone,
        TRUNC(rs.sale_date)        AS last_sale_date,
        rs.store_code              AS last_store
    FROM CUSTOMER c
    JOIN ranked_sales rs ON rs.customer_sid = c.SID AND rs.rn = 1
    LEFT JOIN primary_phone pp ON pp.cust_sid = c.SID AND pp.rn = 1
    WHERE rs.sale_date < TO_DATE(:cutoff_date, 'YYYY-MM-DD')
      AND c.FIRST_NAME IS NOT NULL
      AND UPPER(TRIM(c.FIRST_NAME)) NOT IN ('SYSADMIN', 'TOURIST', 'TOURIST.')
    ORDER BY rs.sale_date DESC, name
"""

# Sheet 2 — customers in CUSTOMER table with NO matching DOCUMENT row.
QUERY_NEVER_PURCHASED = f"""
    WITH {_PRIMARY_PHONE_CTE}
    SELECT
        TRIM(c.FIRST_NAME)  AS name,
        pp.phone_no         AS phone
    FROM CUSTOMER c
    LEFT JOIN primary_phone pp ON pp.cust_sid = c.SID AND pp.rn = 1
    WHERE c.FIRST_NAME IS NOT NULL
      AND UPPER(TRIM(c.FIRST_NAME)) NOT IN ('SYSADMIN', 'TOURIST', 'TOURIST.')
      AND NOT EXISTS (
          SELECT 1
          FROM DOCUMENT d
          WHERE d.BT_CUID = c.SID
            AND d.STATUS = 4
            AND d.receipt_type IN (0, 1)
      )
    ORDER BY name
"""


def _open_connection():
    conn = get_oracle_connection()
    if conn is None:
        raise RuntimeError('Could not open Oracle connection')
    return conn


def fetch_lapsed(cutoff_date: str):
    conn = _open_connection()
    try:
        cur = conn.cursor()
        try:
            cur.execute(QUERY_LAPSED, {'cutoff_date': cutoff_date})
            cols = [d[0].lower() for d in cur.description]
            rows = cur.fetchall()
        finally:
            cur.close()
    finally:
        conn.close()
    return [dict(zip(cols, r)) for r in rows]


def fetch_never_purchased():
    conn = _open_connection()
    try:
        cur = conn.cursor()
        try:
            cur.execute(QUERY_NEVER_PURCHASED)
            cols = [d[0].lower() for d in cur.description]
            rows = cur.fetchall()
        finally:
            cur.close()
    finally:
        conn.close()
    return [dict(zip(cols, r)) for r in rows]


def _coerce_date(value):
    """Oracle DATE -> datetime.date for clean Excel cell formatting."""
    if value is None:
        return None
    if hasattr(value, 'date'):
        return value.date()
    return value


def write_xlsx(lapsed_rows, never_rows, out_path: Path,
               cutoff_date: str, as_of: date):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = ['name', 'phone', 'last_sale_date', 'last_store']
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='305496', end_color='305496', fill_type='solid')

    def _style_header(ws):
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

    def _autosize(ws):
        for col_idx, header in enumerate(headers, start=1):
            cells = [ws.cell(row=r, column=col_idx).value
                     for r in range(2, ws.max_row + 1)] if ws.max_row > 1 else []
            max_len = max(len(str(header)),
                          *(len(str(v or '')) for v in cells)) if cells else len(header)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 14), 40)

    wb = Workbook()

    # ---- Sheet 1: Lapsed ----
    ws1 = wb.active
    ws1.title = 'Lapsed'
    ws1.append(headers)
    _style_header(ws1)
    for r in lapsed_rows:
        ws1.append([
            r.get('name') or '',
            r.get('phone') or '',
            _coerce_date(r.get('last_sale_date')),
            r.get('last_store') or '',
        ])
    for row_idx in range(2, ws1.max_row + 1):
        ws1.cell(row=row_idx, column=3).number_format = 'YYYY-MM-DD'
    ws1.freeze_panes = 'A2'
    _autosize(ws1)

    # ---- Sheet 2: Never Purchased ----
    ws2 = wb.create_sheet('Never Purchased')
    ws2.append(headers)
    _style_header(ws2)
    for r in never_rows:
        ws2.append([
            r.get('name') or '',
            r.get('phone') or '',
            None,   # last_sale_date intentionally blank
            None,   # last_store intentionally blank
        ])
    ws2.freeze_panes = 'A2'
    _autosize(ws2)

    # ---- Metadata ----
    meta = wb.create_sheet('Metadata')
    meta.append(['Parameter', 'Value'])
    for cell in meta[1]:
        cell.font = header_font
        cell.fill = header_fill
    meta.append(['cutoff_date (lapsed: last sale strictly before)', cutoff_date])
    meta.append(['as_of_date (today)', as_of.isoformat()])
    meta.append(['lapsed_count', len(lapsed_rows)])
    meta.append(['never_purchased_count', len(never_rows)])
    meta.append(['name source', 'CUSTOMER.FIRST_NAME (Vietnamese full-name convention)'])
    meta.append(['phone source', 'CUSTOMER_PHONE — PRIMARY_FLAG=1 first, lowest SEQ_NO breaks ties'])
    meta.append(['excluded', 'NULL name, SYSADMIN, Tourist'])
    meta.column_dimensions['A'].width = 44
    meta.column_dimensions['B'].width = 60

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    parser.add_argument('--cutoff-date', default='2025-06-15',
                        help='Sheet 1: customers with last sale strictly before this date (default: 2025-06-15)')
    parser.add_argument('--out-dir', default='document/reports',
                        help='Output directory (default: document/reports)')
    args = parser.parse_args()

    cutoff = args.cutoff_date
    as_of = date.today()

    print(f'[1/3] Querying lapsed customers (last sale < {cutoff}) ...')
    lapsed = fetch_lapsed(cutoff)
    print(f'      -> {len(lapsed)} lapsed customers')

    print('[2/3] Querying never-purchased customers ...')
    never = fetch_never_purchased()
    print(f'      -> {len(never)} never-purchased customers')

    out_path = Path(args.out_dir) / f'dormant_customers_{cutoff}_{as_of.isoformat()}.xlsx'
    print(f'[3/3] Writing {out_path} ...')
    write_xlsx(lapsed, never, out_path, cutoff, as_of)

    print(f'\nDone. Written to:')
    print(f'  {out_path.resolve()}')
    print(f'  Sheet 1 (Lapsed):          {len(lapsed)} rows')
    print(f'  Sheet 2 (Never Purchased): {len(never)} rows')


if __name__ == '__main__':
    main()
