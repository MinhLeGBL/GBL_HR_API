"""Unit tests for app.modules.weekly_report.excel.

Focused on the sheet's column outline: the prior-period block is grouped and
opens collapsed. Everything here goes through a real save/load round trip
rather than inspecting the in-memory worksheet, because the outline only
matters if it survives serialisation into the file Excel opens.
"""
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.modules.weekly_report.excel import (BAD_COLOUR, BAD_ON_DARK,
                                               BRAND_BEIGE, BRAND_DARK,
                                               BRAND_LIGHT, BRAND_ON_DARK,
                                               METRIC_COLUMNS, SHEET_NAMES,
                                               build_workbook_bytes)
from app.modules.weekly_report.period import period_windows

WIDTH = len(METRIC_COLUMNS)
FIRST_CURRENT = 3
FIRST_PRIOR = FIRST_CURRENT + WIDTH
FIRST_DELTA = FIRST_CURRENT + 2 * WIDTH


def _dept_row(day, store, dept):
    return {'day': day, 'store_code': store, 'store_name': f'Store {store}',
            'department': dept, 'sale_net': Decimal('5000000'),
            'return_net': Decimal('100000'), 'sale_gross': Decimal('6000000'),
            'qty_sold': Decimal('3'), 'qty_returned': Decimal('1'),
            'dept_bills': 2}


@pytest.fixture(scope='module')
def workbook():
    """One workbook covering both the current and prior-year windows."""
    days = ([date(2026, 9, d) for d in range(1, 14)] +
            [date(2025, 9, d) for d in range(1, 14)])
    dept_rows = [_dept_row(d, s, dep)
                 for d in days for s in ('S1', 'S2') for dep in ('Alpha', 'Beta')]
    store_rows = [{'day': d, 'store_code': s, 'bills': 5}
                  for d in days for s in ('S1', 'S2')]
    windows = period_windows(date(2026, 9, 14), 'monday', date(2026, 9, 13))
    return load_workbook(BytesIO(
        build_workbook_bytes(windows, dept_rows, store_rows)))


def _cols(first):
    return [get_column_letter(first + i) for i in range(WIDTH)]


@pytest.fixture(params=sorted(SHEET_NAMES.values()))
def sheet(request, workbook):
    """Every period sheet, so a rule cannot hold on one sheet alone."""
    return workbook[request.param]


@pytest.fixture(scope='module')
def real_store_workbook():
    """A workbook using the real store codes, fed in scrambled order."""
    codes = ('RWD', 'RHN', 'HQ', 'RWT', 'RWR', 'RWP')
    days = [date(2026, 9, d) for d in range(7, 14)]
    dept_rows = [_dept_row(d, c, 'RTW') for d in days for c in codes]
    store_rows = [{'day': d, 'store_code': c, 'bills': 5}
                  for d in days for c in codes]
    windows = period_windows(date(2026, 9, 14), 'monday', date(2026, 9, 13))
    return load_workbook(BytesIO(
        build_workbook_bytes(windows, dept_rows, store_rows)))


class TestBrandPalette:
    """The company's own colours, lifted from the Overview deck's P&L sheet.

    Asserted by hex because that is the whole point — "looks about right" is
    how a palette drifts one shade at a time until it no longer matches the
    deck it was copied from.
    """

    def _fill(self, ws, row, col):
        cell = ws.cell(row, col)
        return cell.fill.fgColor.rgb if cell.fill.patternType == 'solid' else None

    def _font_colour(self, ws, row, col):
        c = ws.cell(row, col).font.color
        return c.rgb if (c and isinstance(c.rgb, str)) else None

    def test_the_three_header_blocks_are_the_three_brand_tones(self, sheet):
        # Distinct tones rather than one flat band: the prior block opens
        # COLLAPSED, so the usual view is Current beside Δ with nothing in
        # between, and a single tone would lose that boundary.
        assert self._fill(sheet, 5, FIRST_CURRENT) == f'00{BRAND_DARK}'
        assert self._fill(sheet, 5, FIRST_PRIOR) == f'00{BRAND_LIGHT}'
        assert self._fill(sheet, 5, FIRST_DELTA) == f'00{BRAND_BEIGE}'

    def test_the_metric_row_matches_its_block(self, sheet):
        for col in (FIRST_CURRENT, FIRST_PRIOR, FIRST_DELTA):
            assert self._fill(sheet, 6, col) == self._fill(sheet, 5, col)

    def test_text_on_the_dark_header_is_white(self, sheet):
        assert self._font_colour(sheet, 5, FIRST_CURRENT) == f'00{BRAND_ON_DARK}'

    def test_the_hierarchy_runs_detail_subtotal_grand_total(self, workbook):
        # The part of the deck's scheme worth copying: it uses no borders, so
        # depth is read entirely from fill.
        ws = workbook['WTD']
        grand = next(r for r in range(7, ws.max_row + 1)
                     if ws.cell(r, 1).value == 'ALL STORES')
        total = next(r for r in range(7, grand)
                     if ws.cell(r, 2).value == 'TOTAL')
        detail = next(r for r in range(7, total)
                      if ws.cell(r, 2).value not in (None, 'TOTAL'))

        assert self._fill(ws, detail, FIRST_CURRENT) is None, 'detail = unfilled'
        assert self._fill(ws, total, FIRST_CURRENT) == f'00{BRAND_BEIGE}'
        assert self._fill(ws, grand, FIRST_CURRENT) == f'00{BRAND_DARK}'
        assert self._font_colour(ws, grand, 1) == f'00{BRAND_ON_DARK}'

    def test_the_raw_data_header_wears_the_same_band(self, workbook):
        # One workbook — a reader landing here from a tab should not feel they
        # changed documents.
        raw = workbook['Raw Data']
        assert self._fill(raw, 1, 1) == f'00{BRAND_DARK}'
        assert self._font_colour(raw, 1, 1) == f'00{BRAND_ON_DARK}'


class TestUnfavourableStaysReadable:
    """The red warning has to survive the dark grand-total row."""

    def test_the_dark_row_uses_a_red_that_reads_on_it(self, workbook):
        # C00000 on BRAND_DARK is 1.33:1 — it renders, so nothing fails, and
        # the warning silently disappears from the most-read row on the sheet.
        ws = workbook['WTD']
        grand = next(r for r in range(7, ws.max_row + 1)
                     if ws.cell(r, 1).value == 'ALL STORES')
        colours = set()
        for c in range(FIRST_DELTA, FIRST_DELTA + WIDTH):
            col = ws.cell(grand, c).font.color
            if col and isinstance(col.rgb, str):
                colours.add(col.rgb)
        assert f'00{BAD_COLOUR}' not in colours, 'the dark row must not use C00000'
        assert colours <= {f'00{BAD_ON_DARK}', f'00{BRAND_ON_DARK}'}, colours

    def test_light_rows_keep_the_original_red(self, workbook):
        ws = workbook['WTD']
        found = set()
        for r in range(7, ws.max_row + 1):
            if ws.cell(r, 1).value == 'ALL STORES':
                continue
            for c in range(FIRST_DELTA, FIRST_DELTA + WIDTH):
                col = ws.cell(r, c).font.color
                if col and isinstance(col.rgb, str) and col.rgb.endswith(BAD_COLOUR):
                    found.add(col.rgb)
        assert found, 'the fixture should produce at least one unfavourable move'

    def test_contrast_is_checked_not_eyeballed(self):
        def luminance(hexs):
            ch = [int(hexs[i:i + 2], 16) / 255 for i in (0, 2, 4)]
            ch = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4
                  for x in ch]
            return 0.2126 * ch[0] + 0.7152 * ch[1] + 0.0722 * ch[2]

        def ratio(a, b):
            la, lb = luminance(a), luminance(b)
            return (max(la, lb) + 0.05) / (min(la, lb) + 0.05)

        assert ratio(BAD_ON_DARK, BRAND_DARK) >= 4.0
        assert ratio(BRAND_ON_DARK, BRAND_DARK) >= 4.5
        assert ratio('000000', BRAND_BEIGE) >= 4.5
        assert ratio('000000', BRAND_LIGHT) >= 4.5
        # And the warning must not blur into the ordinary white figures
        # beside it on that same dark row.
        assert ratio(BAD_ON_DARK, BRAND_ON_DARK) >= 1.8


class TestStoreOrderInTheSheet:
    """The sheet and the page must agree, row for row.

    They used to sort independently — `sorted(...)` in three separate places —
    so a reader comparing the attachment against the screen had to take it on
    trust that row three was the same store in both.
    """

    ORDER = ['HQ', 'RWP', 'RHN', 'RWR', 'RWT', 'RWD']

    def _store_rows_in_order(self, ws):
        seen = []
        for row in range(7, ws.max_row + 1):
            code = ws.cell(row, 1).value
            if code and code not in seen and code != 'ALL STORES':
                seen.append(code)
        return seen

    def test_rows_follow_the_boards_order(self, real_store_workbook):
        for name in SHEET_NAMES.values():
            ws = real_store_workbook[name]
            assert self._store_rows_in_order(ws) == self.ORDER, name

    def test_it_is_not_alphabetical(self, real_store_workbook):
        order = self._store_rows_in_order(real_store_workbook['WTD'])
        assert order != sorted(order)

    def test_the_legend_follows_the_same_order(self, real_store_workbook):
        # A legend in a different order to the rows is worse than none: the
        # reader runs down it looking for the store they are on and finds it
        # somewhere else.
        legend = real_store_workbook['WTD']['A4'].value
        positions = [legend.index(f'{c} =') for c in self.ORDER]
        assert positions == sorted(positions), legend


class TestWhichSheetsTheWorkbookCarries:
    """Which periods get exported is a decision, not an implementation detail.

    The `sheet` fixture above parametrises over SHEET_NAMES, so it follows any
    edit to that dict silently — swapping a period in or out would not fail a
    single test there. These assert the decision itself.
    """

    def test_the_workbook_carries_wtd_mtd_and_ytd(self, workbook):
        # 'Raw Data' trails the period sheets and is not one of them.
        assert workbook.sheetnames[:3] == ['WTD', 'MTD', 'YTD']
        assert set(workbook.sheetnames) - {'Raw Data'} == set(SHEET_NAMES.values())

    def test_week_by_week_is_not_exported(self, workbook):
        # Page-only since 2026-09-22: readers asked for the year-on-year week
        # in the attachment instead. Re-adding it is a decision, so it should
        # break this test rather than pass unnoticed.
        assert 'Week by Week' not in workbook.sheetnames

    def test_every_sheet_compares_against_a_year_earlier(self, workbook):
        # The point of the swap: one baseline across the whole workbook, so a
        # reader moving between sheets never has to check which one they are on.
        for name in SHEET_NAMES.values():
            ws = workbook[name]
            headers = [ws.cell(5, c).value for c in range(1, 30)]
            assert any(h and 'Prior year' in str(h) for h in headers), name


class TestWtdSheetHeader:
    """WTD compares whole ISO weeks a year apart, and says so."""

    def test_the_title_names_the_iso_week_on_both_sides(self, workbook):
        # Two date ranges a year apart leave the reader to verify they really
        # are the same week; the week numbers state it.
        title = workbook['WTD']['A1'].value
        assert 'Week 37 2026' in title and 'Week 37 2025' in title

    def test_the_note_warns_that_the_dates_do_not_line_up(self, workbook):
        # Week 37 of 2026 is 07/09-13/09; of 2025, 08/09-14/09. A reader who
        # spots the one-day offset should find it explained, not suspect a bug.
        note = workbook['WTD']['A3'].value
        assert 'SAME ISO week' in note
        assert 'differ by a day or two' in note

    def test_the_note_names_this_sheets_own_week_number(self, workbook):
        # It used to carry a hardcoded "week 38 against week 38", which read as
        # a contradiction on every sheet that was not week 38 — the fastest way
        # to make a reader distrust a note that is otherwise correct.
        assert 'week 37 against week 37' in workbook['WTD']['A3'].value
        assert 'week 38' not in workbook['WTD']['A3'].value

    def test_the_blocks_carry_the_week_and_its_dates(self, workbook):
        ws = workbook['WTD']
        current = ws.cell(5, FIRST_CURRENT).value
        prior = ws.cell(5, FIRST_PRIOR).value
        assert 'Week 37 2026' in current and '2026-09-07' in current
        assert 'Week 37 2025' in prior and '2025-09-08' in prior


class TestPriorBlockCollapsed:
    def test_prior_columns_are_grouped_at_outline_level_1(self, sheet):
        levels = {sheet.column_dimensions[c].outlineLevel for c in _cols(FIRST_PRIOR)}
        assert levels == {1}

    def test_prior_columns_open_hidden(self, sheet):
        # "Collapsed by default" means the file opens with them already hidden,
        # not merely groupable.
        assert all(sheet.column_dimensions[c].hidden for c in _cols(FIRST_PRIOR))

    def test_the_summary_column_carries_the_collapsed_flag(self, sheet):
        # Excel draws the +/- control on the column AFTER the group and marks
        # that one collapsed. Without this the columns are hidden but the
        # button does not render as expandable.
        summary = get_column_letter(FIRST_DELTA)
        assert sheet.column_dimensions[summary].collapsed is True

    def test_summary_is_to_the_right(self, sheet):
        # summaryRight is openpyxl's default, but the button landing on the
        # delta block is load-bearing rather than incidental.
        assert sheet.sheet_properties.outlinePr.summaryRight is True

    def test_current_and_delta_blocks_stay_visible(self, sheet):
        for first in (FIRST_CURRENT, FIRST_DELTA):
            assert not any(sheet.column_dimensions[c].hidden for c in _cols(first))

    def test_store_and_department_columns_stay_visible(self, sheet):
        assert not sheet.column_dimensions['A'].hidden
        assert not sheet.column_dimensions['B'].hidden


class TestPriorWidthsSurviveGrouping:
    """`column_dimensions.group()` would have destroyed these.

    It deletes the per-column dimensions across the range and leaves one
    spanning dimension keyed to the first column, so the rest fall back to
    Excel's default 13 — measured: group('J','P') over [10,8,7,9,9,8,9] gives
    [10,13,13,13,13,13,13]. The block would then be wider on expand than the
    layout was sized for. The outline attributes are therefore set per column.
    """

    def test_each_prior_column_keeps_its_own_width(self, sheet):
        prior = [sheet.column_dimensions[c].width for c in _cols(FIRST_PRIOR)]
        current = [sheet.column_dimensions[c].width for c in _cols(FIRST_CURRENT)]
        assert prior == current

    def test_widths_are_not_all_equal(self, sheet):
        # The regression this guards against collapses them to one value, so a
        # per-metric spread is the signal that it has not happened.
        assert len({sheet.column_dimensions[c].width
                    for c in _cols(FIRST_PRIOR)}) > 1

    def test_no_prior_column_fell_back_to_the_excel_default(self, sheet):
        assert all(sheet.column_dimensions[c].width != 13
                   for c in _cols(FIRST_PRIOR))


class TestRowOutlineStillWorks:
    """The column outline shares `outlinePr` with the existing row outline."""

    def test_summary_below_is_preserved(self, sheet):
        # Department rows sit ABOVE their store TOTAL, so this must stay True
        # or the row +/- control detaches from the TOTAL row.
        assert sheet.sheet_properties.outlinePr.summaryBelow is True

    def test_department_rows_are_still_grouped_and_hidden(self, sheet):
        grouped = [r for r, d in sheet.row_dimensions.items()
                   if d.outlineLevel == 1]
        assert grouped, 'department rows lost their outline level'
        assert all(sheet.row_dimensions[r].hidden for r in grouped)

    def test_some_total_row_is_marked_collapsed(self, sheet):
        assert any(d.collapsed for d in sheet.row_dimensions.values())


class TestRawDataSheetIsUntouched:
    def test_raw_sheet_has_no_hidden_columns(self, workbook):
        ws = workbook['Raw Data']
        assert not any(d.hidden for d in ws.column_dimensions.values())
